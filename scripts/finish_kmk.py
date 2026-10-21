from pathlib import Path
import csv, json, re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
B=Path(__file__).resolve().parents[1]
p=B/'data/KANDIDAT_SBM_SBK_PERLU_VERIFIKASI.csv'
with p.open(encoding='utf-8-sig',newline='') as f: data=list(csv.DictReader(f))
# Excel reviewer
w=Workbook();s=w.active;s.title='Review Indeks KMK';cols=list(data[0]);s.append(cols)
for i,r in enumerate(data,2):
 s.append([int(r[c]) if c in ('halaman_pdf','besaran_rupiah') else r[c] for c in cols])
 for c in s[i]:
  c.alignment=Alignment(vertical='top',wrap_text=c.column in [8,9,14,16])
s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
for c in s[1]:c.font=Font(color='FFFFFF',bold=True);c.fill=PatternFill('solid',fgColor='174A78');c.alignment=Alignment(wrap_text=True)
widths=[23,10,19,10,14,13,18,22,58,15,12,18,21,65,22,52,36]
for j,wd in enumerate(widths,1):s.column_dimensions[get_column_letter(j)].width=wd
for row in s.iter_rows(min_row=2): row[12].number_format='#,##0';row[14].fill=PatternFill('solid',fgColor='FFF0D6')
info=w.create_sheet('Petunjuk')
for rr in [ ['PANDUAN REVIEW KANDIDAT KMK 296/2026 dan KMK 57/2026'], ['Seluruh baris bersifat PERLU_REVIEW dan bukan publikasi angka final.'],['Periksa setiap baris terhadap PDF asli pada halaman_pdf, termasuk kode, kelompok, satuan, besaran, dan wilayah.'],['Ganti status_verifikasi menjadi DISETUJUI hanya setelah seluruh kolom sesuai; isi catatan verifikasi.'],['Data di workbook tidak otomatis diimpor ke D1; gunakan skrip export_approved_sql.py setelah review.'],['SBM: KMK 296/2026 digunakan untuk TA 2027.'],['SBK: KMK 57/2026 menyebut penyusunan RKA-K/L mulai 2026.']]:info.append(rr)
info.column_dimensions['A'].width=110
w.save(B/'data/REVIEW_KMK_296_57.xlsx')
# D1 staging tables (independent of existing Google sync schema)
(B/'migrations/0002_kmk_review.sql').write_text('''-- Safe addition: separate KMK tables; existing publications, items and rates remain unchanged.
CREATE TABLE IF NOT EXISTS kmk_review_candidates (
  id TEXT PRIMARY KEY, jenis TEXT NOT NULL, regulasi TEXT NOT NULL,
  lampiran TEXT, halaman_pdf INTEGER NOT NULL, kode_item TEXT, kode_program TEXT,
  kelompok TEXT, uraian TEXT NOT NULL, wilayah TEXT, volume TEXT, satuan TEXT,
  besaran_rupiah INTEGER NOT NULL, teks_asli TEXT,
  status_verifikasi TEXT NOT NULL DEFAULT 'PERLU_REVIEW'
    CHECK (status_verifikasi IN ('PERLU_REVIEW','DISETUJUI','DITOLAK')),
  catatan_verifikasi TEXT, sumber_pdf TEXT NOT NULL
);
CREATE INDEX IF NOT EXISTS idx_kmk_review_search ON kmk_review_candidates (jenis, regulasi, halaman_pdf);
-- Populate only with explicitly approved and verified rows via scripts/export_approved_sql.py.
''',encoding='utf-8')
(B/'scripts/export_approved_sql.py').write_text('''#!/usr/bin/env python3
"""Export verified CSV rows to SQL for separate D1 staging table. Never modifies production catalog."""
import csv,sys,pathlib
if len(sys.argv)<3: raise SystemExit('Usage: python scripts/export_approved_sql.py REVIEWED.csv approved.sql')
q=lambda s: "'"+str(s).replace("'","''")+"'"
with open(sys.argv[1],encoding='utf-8-sig',newline='') as f: rows=list(csv.DictReader(f))
cols=['id','jenis','regulasi','lampiran','halaman_pdf','kode_item','kode_program','kelompok','uraian','wilayah','volume','satuan','besaran_rupiah','teks_asli','status_verifikasi','catatan_verifikasi','sumber_pdf']
with open(sys.argv[2],'w',encoding='utf8') as f:
 f.write('-- Reviewed staging data. Apply after migrations/0002_kmk_review.sql.\\nBEGIN TRANSACTION;\\n')
 count=0
 for r in rows:
  if r['status_verifikasi']!='DISETUJUI':continue
  if not r['catatan_verifikasi'].strip():raise ValueError('Catatan verifikasi wajib untuk '+r['id'])
  if not r['uraian'].strip() or not int(r['besaran_rupiah'])>0: raise ValueError('Kandidat tidak lengkap '+r['id'])
  vals=','.join(str(int(r[c])) if c in ('halaman_pdf','besaran_rupiah') else q(r[c]) for c in cols)
  f.write('INSERT OR REPLACE INTO kmk_review_candidates ('+','.join(cols)+') VALUES ('+vals+');\\n')
  count+=1
 f.write('COMMIT;\\n')
print('Exported approved rows:',count)
''',encoding='utf-8')
app=B/'src/app.js';t=app.read_text()
t=t.replace("pustaka:'Pustaka',faq:'FAQ'","pustaka:'Pustaka',indeks:'Indeks KMK',faq:'FAQ'")
new='''
let biayaCandidates=null;
function indeksKmk(){return `<section class="page-head"><span class="small-label">EKSTRAKSI TERSTRUKTUR DOKUMEN</span><h1>Indeks Besaran SBM & SBK</h1><p>Cari uraian, nominal rupiah, satuan, kode SBK, atau nomor regulasi dari PDF sumber.</p></section><div class="reg-advisory"><b>PENTING:</b> Semua data pada halaman ini adalah hasil pembacaan otomatis dan berstatus <b>PERLU REVIEW</b>. Periksa lampiran PDF sebelum menggunakannya dalam perencanaan anggaran. Data tidak otomatis ditambahkan ke D1 resmi.</div><form id="biaya-query-form" class="toolbar"><label class="filter-search">${icon('search',20)}<input id="biaya-q" placeholder="Cari: honorarium, uang makan, layanan audit, 630000..." aria-label="Cari kandidat besaran biaya" /></label><select id="biaya-kind" aria-label="Jenis standar biaya"><option value="">Semua jenis</option><option value="SBM">SBM (KMK 296)</option><option value="SBK">SBK (KMK 57)</option></select><button type="submit" class="secondary">Cari indeks</button></form><div class="result-note">Hasil tampil maksimum 100 baris. Klik sumber untuk mengecek angka pada halaman PDF.</div><div id="biaya-result"><p class="empty">Masukkan kata kunci untuk melihat kandidat data.</p></div>`;}
async function lookupBiaya(e){e.preventDefault();const target=document.querySelector('#biaya-result');const q=document.querySelector('#biaya-q').value.trim().toLowerCase();const kind=document.querySelector('#biaya-kind').value;if(q.length<2){target.innerHTML='<p class="empty">Masukkan paling sedikit dua karakter.</p>';return;}target.textContent='Memuat indeks terstruktur...';try{if(!biayaCandidates){let resp=await fetch('/data/kandidat-indeks.json',{credentials:'same-origin',cache:'no-store'});if(!resp.ok)throw Error('File indeks belum tersedia');biayaCandidates=await resp.json();}const results=biayaCandidates.rows.filter(r=>(!kind||r.jenis===kind)&&[r.uraian,r.kode_item,r.kode_program,r.satuan,r.besaran_rupiah,r.regulasi,r.halaman_pdf].join(' ').toLowerCase().includes(q));const found=results.slice(0,100);target.innerHTML=`<p class="result-note">Ditemukan ${results.length} kandidat. Maksimum 100 ditampilkan; semuanya belum diverifikasi.</p>`+(found.length?`<div class="kmk-index-list">${found.map(r=>`<article class="kmk-index-card"><div class="entry-head">${badge(r.jenis)} ${badge(r.regulasi)} <span class="verification pending">PERLU REVIEW</span></div><h3>${escapeHTML(r.uraian)}</h3><div><strong>Rp ${Number(r.besaran_rupiah).toLocaleString('id-ID')}</strong> ${escapeHTML(r.volume?r.volume+' ':'')}${escapeHTML(r.satuan)}</div><p>${escapeHTML(r.kode_program||r.kode_item||'')} · Lampiran ${escapeHTML(r.lampiran)} · Halaman PDF ${r.halaman_pdf}</p><a href="${escapeHTML(r.sumber_pdf)}#page=${r.halaman_pdf}" target="_blank" rel="noopener noreferrer">Periksa halaman sumber ↗</a></article>`).join('')}</div>`:'<p class="empty">Tidak ada kandidat cocok. Gunakan pencarian teks PDF pada menu Regulasi.</p>');}catch(e){target.innerHTML='<p class="empty">Indeks tidak tersedia. Pastikan file /data/kandidat-indeks.json ikut terunggah.</p>';}}
'''
t=t.replace('const content=()=>',new+'\nconst content=()=>')
t=t.replace("state.page==='pustaka'?simple('Dokumen','Pustaka Standar Biaya','Dokumen referensi yang telah disetujui untuk dipublikasikan.',state.docs):faq()","state.page==='pustaka'?simple('Dokumen','Pustaka Standar Biaya','Dokumen referensi yang telah disetujui untuk dipublikasikan.',state.docs):state.page==='indeks'?indeksKmk():faq()")
t=t.replace("document.querySelector('#kmk-search-form')?.addEventListener('submit',kmkLookup);","document.querySelector('#kmk-search-form')?.addEventListener('submit',kmkLookup);document.querySelector('#biaya-query-form')?.addEventListener('submit',lookupBiaya);")
app.write_text(t)
with (B/'src/styles.css').open('a') as f:f.write('''\n/* Searchable candidate KMK entries, separated from verified D1 catalog */\n.kmk-index-list{display:grid;grid-template-columns:repeat(auto-fit,minmax(min(100%,310px),1fr));gap:15px;margin:20px 0}.kmk-index-card{background:#fff;border:1px solid #d6e9f5;border-radius:15px;padding:19px;box-shadow:0 4px 18px rgba(30,88,139,.05)}.kmk-index-card h3{font-size:1rem;color:#15466e;margin:12px 0}.kmk-index-card strong{font-size:1.1rem;color:#075b99}.kmk-index-card p{color:#57718a;font-size:.88rem}.kmk-index-card a{font-weight:700;color:#0769ae}#biaya-query-form{display:flex;gap:12px;flex-wrap:wrap}#biaya-query-form .filter-search{flex:1;min-width:240px}#biaya-query-form select{padding:12px;border:1px solid #cde0ed;border-radius:12px;background:#fff}\n''')
(B/'docs/PANDUAN-DATA-TERSTRUKTUR.md').write_text('''# Indeks terstruktur KMK 296/2026 dan KMK 57/2026

## Isi
- Menu **Indeks KMK**: pencarian baris kandidat di browser, tautan halaman asli PDF.
- `data/KANDIDAT_SBM_SBK_PERLU_VERIFIKASI.csv`: hasil ekstraksi otomatis.
- `data/REVIEW_KMK_296_57.xlsx`: lembar kerja review.
- `data/review_kmk_296_57.sqlite`: basis data lokal untuk pemeriksaan, **bukan** DB produksi.
- `public/data/kandidat-indeks.json`: dataset indeks web, berstatus PERLU_REVIEW.
- `migrations/0002_kmk_review.sql`: tabel D1 terpisah, tidak mengubah katalog D1 yang ada.
- `scripts/export_approved_sql.py`: menghasilkan SQL hanya untuk baris yang diverifikasi dan berstatus DISETUJUI.

## Cakupan dan keterbatasan
Ekstraksi teks PDF bukan pembacaan tabel yang sepenuhnya akurat. Bisa ada deskripsi terpotong, kode tidak terbaca, angka salah terhubung ke baris, pengulangan, dan baris tidak terekstraksi. Semua 5.715 baris kandidat **wajib diperiksa manual** sebelum disahkan. Jangan gunakan angka kandidat sebagai referensi final. Jumlah baris bukan jumlah resmi item dalam lampiran.

## Proses aman
1. Periksa halaman sumber PDF dan cocokkan setiap kandidat (nomenklatur, kelompok, kode, besaran, satuan, wilayah dan lampiran).
2. Isi bagian yang kosong, perbaiki data bila perlu, beri `status_verifikasi=DISETUJUI` serta catatan pengecekan, dan simpan CSV UTF-8.
3. `python scripts/export_approved_sql.py data/REVIEWED.csv data/approved.sql`.
4. Uji migrasi dan SQL hasil review **di database pengujian terlebih dahulu**, bukan produksi.
5. Tabel `kmk_review_candidates` sengaja terpisah dari `items`, `rates`, dan mekanisme publikasi Google Sheets. Tidak ada auto-publish.

## Catatan regulasi
- KMK 296/2026 menyebut besaran SBM digunakan untuk TA 2027.
- KMK 57/2026 menyebut indeks SBK digunakan untuk penyusunan RKA-K/L mulai tahun 2026.
''',encoding='utf-8')
(B/'README.md').write_text((B/'README.md').read_text()+ '\n\n## Indeks Data Terstruktur\nLihat `docs/PANDUAN-DATA-TERSTRUKTUR.md`. Data hasil ekstraksi PDF berstatus PERLU_REVIEW dan tidak otomatis ditulis ke database operasional.\n')
print('finished',len(data))
